#!/usr/bin/env python3
"""生成「闭源 API vs 开源权重发布」对比 PPT（阿里巴巴配色）。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# 阿里配色
ORANGE   = RGBColor(0xFF, 0x6A, 0x00)
ORANGE_D = RGBColor(0xE6, 0x4A, 0x19)
TEAL     = RGBColor(0x0B, 0xA5, 0xA5)
INK      = RGBColor(0x1F, 0x23, 0x29)
INK_DIM  = RGBColor(0x5B, 0x64, 0x70)
TEXT     = RGBColor(0x34, 0x40, 0x4D)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT  = RGBColor(0xFA, 0xFA, 0xFA)
LINE     = RGBColor(0xEC, 0xEC, 0xEF)

ROWS = [
    ("模型能力",   "同名版本通常更强，持续优化迭代",            "发布时刻的固定快照，能力不再更新"),
    ("上下文长度", "更长（如 Qwen-Long 1M tokens）",            "通常 32K/128K，需自行扩展"),
    ("训练数据",   "可能用了更多 / 更新的数据",                  "发布时的训练数据截止"),
    ("对齐与安全", "更严格，根据线上反馈持续调整",              "固定状态，用户可微调修改"),
    ("推理性能",   "专有推理加速，低延迟高吞吐",                "自行部署优化（vLLM / TGI 等）"),
    ("功能完整度", "Tool calling、structured output 等更成熟",   "基础能力具备，工程化需自建"),
    ("更新方式",   "静默升级，同 endpoint 背后模型可能已迭代",  "版本固定，等下一个 release"),
    ("数据安全",   "数据经过服务商",                            "私有部署，数据不出域"),
    ("定制化",     "仅支持 API 参数调整",                        "可全量微调 / LoRA / 量化等"),
    ("成本模式",   "按 token 付费，无需 GPU",                    "需自备算力，一次性部署成本高"),
    ("适用场景",   "追求最优效果、快速上线、不敏感数据",        "私有化部署、领域微调、数据合规要求高"),
]

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式

SW, SH = prs.slide_width, prs.slide_height

def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

# 顶部品牌条
from pptx.enum.shapes import MSO_SHAPE
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Pt(7))
set_fill(bar, ORANGE)

# 标题
tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.8))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
for txt, col in [("闭源（API）", ORANGE), (" vs ", INK), ("开源（权重发布）", TEAL)]:
    r = p.add_run(); r.text = txt
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = col
    r.font.name = "PingFang SC"

# 副标题
sb = slide.shapes.add_textbox(Inches(0.57), Inches(1.05), Inches(8), Inches(0.4))
sp = sb.text_frame.paragraphs[0]
sr = sp.add_run(); sr.text = "大模型选型对比 · 11 个关键维度"
sr.font.size = Pt(13); sr.font.color.rgb = INK_DIM; sr.font.name = "PingFang SC"

# 表格
left, top = Inches(0.55), Inches(1.6)
width, height = Inches(12.23), Inches(5.55)
rows, cols = len(ROWS) + 1, 3
gtbl = slide.shapes.add_table(rows, cols, left, top, width, height)
tbl = gtbl.table
tbl.columns[0].width = Inches(2.3)
tbl.columns[1].width = Inches(4.96)
tbl.columns[2].width = Inches(4.97)

# 关闭内置样式条带
tbl.first_row = True
tbl.horz_banding = False

hdr = [("对比维度", INK), ("闭源", ORANGE), ("开源", TEAL)]
for c, (label, color) in enumerate(hdr):
    cell = tbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = color
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.18); cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = label
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    r.font.name = "PingFang SC"
    sub = {1: "API · 托管调用", 2: "权重发布 · 自行部署"}.get(c)
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(9.5); r2.font.color.rgb = WHITE; r2.font.name = "PingFang SC"

for i, (dim, api, oss) in enumerate(ROWS, start=1):
    bg = WHITE if i % 2 else ROW_ALT
    for c, (txt, bold, color) in enumerate([
        (dim, True, INK), (api, False, TEXT), (oss, False, TEXT)
    ]):
        cell = tbl.cell(i, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.18); cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = txt
        r.font.size = Pt(13); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = "PingFang SC"

# 行高
tbl.rows[0].height = Inches(0.62)
for i in range(1, rows):
    tbl.rows[i].height = Inches(0.448)

prs.save("/home/user/gh-contributor/闭源vs开源对比.pptx")
print("saved")
